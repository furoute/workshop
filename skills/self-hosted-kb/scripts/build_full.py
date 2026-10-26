from __future__ import annotations

"""Build entity cards from an inventory and extract source evidence.

Usage:
    python scripts/build_full.py \
        --source-dir <corpus-root> \
        --output-dir <build-output>

The output directory must contain ``entity-inventory/实体清单.csv`` and
``entity-inventory/主体文件台账.csv``. All paths are supplied on the command
line; no machine-specific default is used.
"""

import argparse
import csv
import html
import json
import os
import re
import unicodedata
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation


SOURCE_ROOT: Path
OUTPUT_DIR: Path
VAULT: Path
INVENTORY: Path
BUILD: Path
BATCH_DIR: Path
CACHE_DIR: Path
RUN_DATE: str
WORKERS = 1

TYPE_INFO = {
    "project": ("项目", "项目卡", "项目类实体"),
    "opportunity": ("机会", "机会卡", "机会类实体"),
    "solution": ("方案", "方案卡", "方案类实体"),
    "standard": ("规范", "规范卡", "规范类实体"),
    "procurement": ("采购", "采购卡", "采购类实体"),
}

SAMPLE_CARDS: dict[str, Path] = {}

BATCHES = {
    "1": (("project",), "项目：34张"),
    "2": (("opportunity",), "机会：26张"),
    "3": (("solution",), "方案：27张"),
    "4": (("standard", "procurement"), "规范38张 + 采购4张"),
}

TEXT_EXTENSIONS = {".docx", ".pdf", ".xlsx", ".pptx", ".txt", ".md", ".csv", ".html"}
UNPARSED_EXTENSIONS = {".doc", ".xls", ".ppt", ".wps", ".rtf", ".zip", ".rar", ".7z"}

FACT_TERMS = re.compile(
    r"项目名称|采购项目名称|采购人|采购单位|客户单位|建设单位|使用单位|"
    r"项目编号|采购编号|预算|限价|控制价|报价|成交|中标|合同金额|"
    r"服务期|交付期|建设期|工期|采购内容|服务内容|建设内容|产品清单|"
    r"适用|定位|能力|流程|管理办法|规范|标准|金额"
)
BAD_SENTENCE = re.compile(
    r"建议此部分按照模板|按照模板字眼|此处填写|请输入|待补充|XX年XXX|"
    r"XXX|某某|示例文本|目录$",
    re.IGNORECASE,
)
PARTY_RE = re.compile(
    r"(采购人|采购单位|客户单位|建设单位|使用单位|甲方|"
    r"成交供应商|中标供应商|中标人|投标人|响应人|供应商|"
    r"实施单位|施工单位|乙方)\s*[:：]\s*([^\n|。；]{2,100})"
)
PROJECT_NAME_RE = re.compile(r"(?:项目名称|采购项目名称)\s*[:：]\s*([^\n|。；]{3,100})")
DEFINED_PARTY_RE = re.compile(
    r"([\u4e00-\u9fffA-Za-z0-9（）()·、]{4,80})\s*"
    r"[（(]\s*以下简称\s*[“\"']?(采购人|客户单位|甲方|乙方)[”\"']?\s*[）)]"
)
CUSTOMER_CONTEXT_RE = re.compile(
    r"([\u4e00-\u9fffA-Za-z0-9（）()·]{4,70}"
    r"(?:有限公司|有限责任公司|集团|中心|局|厅|大学|学校|银行|研究院|研究所))"
    r"[^。；\n]{0,80}采购人(?:启动|拟|作为)"
)
AMOUNT_RE = re.compile(
    r"(?P<prefix>[\s\S]{0,140}?)(?P<currency>人民币|RMB|CNY|[¥￥])?\s*"
    r"(?P<num>\d[\d,]*(?:\.\d+)?)\s*(?P<unit>亿元|万元|亿|万|元)"
)
DATE_RE = re.compile(r"20\d{2}[-年.]\d{1,2}[-月.]\d{1,2}日?")


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def write_text(path: Path, value: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(value, encoding="utf-8", newline="\n")


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def clean_text(value: str) -> str:
    value = unicodedata.normalize("NFKC", value or "")
    value = value.replace("\u00a0", " ").replace("\u3000", " ")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def safe_name(value: str, limit: int = 105) -> str:
    value = unicodedata.normalize("NFKC", value)
    value = re.sub(r'[\\/:*?"<>|\r\n]+', "-", value)
    value = re.sub(r"\s+", " ", value).strip(" .-")
    if len(value) > limit:
        value = value[:limit].rstrip(" .-")
    return value


def find_card(entity: dict[str, str]) -> Path:
    eid = entity["entity_id"]
    if eid in SAMPLE_CARDS and SAMPLE_CARDS[eid].exists():
        return SAMPLE_CARDS[eid]
    folder, prefix, _ = TYPE_INFO[entity["entity_type"]]
    return VAULT / folder / f"{prefix}-{safe_name(entity['entity_name'])}.md"


def extract_docx(path: Path) -> str:
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_pptx(path: Path) -> str:
    presentation = Presentation(path)
    parts: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text)
    return "\n".join(parts)


def extract_xlsx(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in workbook.worksheets:
        parts.append(f"SHEET {sheet.title}")
        for row in sheet.iter_rows(
            min_row=1, max_row=min(sheet.max_row or 0, 600), values_only=True
        ):
            values = ["" if value is None else str(value) for value in row]
            if any(values):
                parts.append(" | ".join(values))
    return "\n".join(parts)


def extract_pdf(path: Path) -> str:
    with path.open("rb") as stream:
        reader = PdfReader(stream, strict=False)
        pages = reader.pages[:40]
        return "\n".join((page.extract_text() or "") for page in pages)


def extract_text(path: Path) -> tuple[str, str]:
    extension = path.suffix.casefold()
    if extension not in TEXT_EXTENSIONS:
        return "", "格式未解析"
    try:
        if extension == ".docx":
            text = extract_docx(path)
        elif extension == ".pptx":
            text = extract_pptx(path)
        elif extension == ".xlsx":
            text = extract_xlsx(path)
        elif extension == ".pdf":
            text = extract_pdf(path)
        elif extension == ".html":
            text = html.unescape(re.sub(r"<[^>]+>", " ", path.read_text("utf-8", errors="ignore")))
        else:
            text = path.read_text("utf-8", errors="ignore")
    except Exception as exc:  # noqa: BLE001 - exact extraction failures belong in the QA ledger
        return "", f"解析失败：{type(exc).__name__}: {exc}"
    text = clean_text(text)
    if not text:
        return "", "无可提取正文"
    return text[:300_000], ""


def source_score(row: dict[str, str], entity_type: str, entity_name: str) -> float:
    filename = row["filename"]
    extension = row["extension"].casefold()
    score = 0.0
    if extension == ".docx":
        score += 8
    elif extension == ".pdf":
        score += 7
    elif extension == ".xlsx":
        score += 5
    elif extension == ".pptx":
        score += 4
    elif extension in {".txt", ".md", ".csv", ".html"}:
        score += 3
    else:
        return -1000
    for term, points in {
        "采购文件": 20,
        "招标文件": 20,
        "询比": 18,
        "结果公告": 24,
        "成交": 22,
        "中标": 22,
        "通知书": 20,
        "候选人": 16,
        "合同": 18,
        "技术规范书": 16,
        "技术部分": 11,
        "需求": 10,
        "立项": 10,
        "方案": 8,
        "报价": 8,
        "清单": 6,
        "操作手册": 12,
        "管理办法": 15,
        "模板": -8,
        "副本": -12,
        "模板说明": -8,
    }.items():
        if term in filename:
            score += points
    if "测试" in filename and entity_type != "project":
        score -= 5
    if entity_type == "solution":
        for term, points in {
            "介绍": 8,
            "方案": 10,
            "白皮书": 12,
            "产品": 10,
            "报价": 8,
            "案例": 8,
        }.items():
            if term in filename:
                score += points
    elif entity_type == "standard":
        for term, points in {
            "流程": 10,
            "规范": 12,
            "手册": 12,
            "模型": 10,
            "评估": 8,
            "办法": 12,
        }.items():
            if term in filename:
                score += points
    entity_tokens = [
        token
        for token in re.split(r"[\s\-—_（）()]+", entity_name)
        if len(token) >= 4
    ]
    score += 2 * sum(1 for token in entity_tokens if token in filename)
    try:
        size = int(row["size"])
        if 8_000 <= size <= 20_000_000:
            score += 2
    except (TypeError, ValueError):
        pass
    return score


def sentence_candidates(text: str) -> list[str]:
    rows: list[str] = []
    buffer = ""
    for part in text.splitlines():
        part = re.sub(r"^SHEET\s+", "", part.strip())
        part = re.sub(r"\s+", " ", part)
        if not part:
            continue
        if "|" in part:
            if buffer:
                rows.append(buffer)
                buffer = ""
            if 12 <= len(part) <= 300 and not BAD_SENTENCE.search(part):
                rows.append(part)
            continue
        buffer = f"{buffer}{part}" if buffer else part
        if part[-1:] in "。！？；" or len(buffer) >= 240:
            if 12 <= len(buffer) <= 300 and not BAD_SENTENCE.search(buffer):
                rows.append(buffer)
            buffer = ""
    if buffer and 12 <= len(buffer) <= 300 and not BAD_SENTENCE.search(buffer):
        rows.append(buffer)
    return rows


def extract_facts(text: str, entity_name: str, limit: int = 2) -> list[str]:
    tokens = [
        token
        for token in re.split(r"[\s\-—_（）()]+", entity_name)
        if len(token) >= 3
    ]
    ranked: list[tuple[float, int, str]] = []
    for index, sentence in enumerate(sentence_candidates(text)):
        score = 0.0
        for match in FACT_TERMS.findall(sentence):
            score += 2.0
            if match in {"项目名称", "采购人", "客户单位", "预算", "限价", "成交", "中标", "合同金额"}:
                score += 2.0
        score += sum(1.0 for token in tokens if token in sentence)
        if DATE_RE.search(sentence):
            score += 1.0
        if "模板" in sentence:
            score -= 6.0
        if 25 <= len(sentence) <= 190:
            score += 1.0
        ranked.append((-score, index, sentence))
    ranked.sort()
    result: list[str] = []
    seen: set[str] = set()
    for _, _, sentence in ranked:
        key = re.sub(r"\W+", "", sentence)[:100]
        if key in seen:
            continue
        seen.add(key)
        result.append(sentence if len(sentence) <= 240 else sentence[:237] + "…")
        if len(result) >= limit:
            break
    return result


def party_mentions(text: str) -> list[tuple[str, str]]:
    result: list[tuple[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for role, raw_value in PARTY_RE.findall(text):
        value = re.split(r"\s{2,}|\||；", raw_value, maxsplit=1)[0].strip(" ，,;；")
        if len(value) < 2 or len(value) > 70:
            continue
        if (
            BAD_SENTENCE.search(value)
            or value in {"无", "待定", "详见采购文件"}
            or re.search(
                r"[\[\]（）()]|加盖|公章|签字|承诺|所有|任何|其他|响应供应商|本项目|该供应商",
                value,
            )
            or not re.search(
                r"公司|集团|中心|局|厅|院|大学|学校|银行|研究所|委员会|"
                r"事业部|分公司|研究院|管理处|厂",
                value,
            )
        ):
            continue
        key = (role, value)
        if key not in seen:
            seen.add(key)
            result.append(key)
    for value, role in DEFINED_PARTY_RE.findall(text):
        value = value.strip(" ，,;；")
        if (
            len(value) >= 4
            and not re.search(r"[\[\]（）()]|加盖|公章|签字|承诺", value)
            and re.search(
                r"公司|集团|中心|局|厅|院|大学|学校|银行|研究所|委员会|"
                r"事业部|分公司|研究院|管理处|厂",
                value,
            )
        ):
            key = (role, value)
            if key not in seen:
                seen.add(key)
                result.append(key)
    for value in CUSTOMER_CONTEXT_RE.findall(text):
        value = re.sub(
            r"^(?:深入)?(?:贯彻)?落实|^(?:面向|服务|依托|根据|为|由|向|对)",
            "",
            value,
        ).strip()
        key = ("客户单位", value)
        if key not in seen:
            seen.add(key)
            result.append(key)
    return result


def normalize_amount(number: str, unit: str) -> float:
    value = float(number.replace(",", ""))
    if unit in {"亿元", "亿"}:
        return value * 100_000_000
    if unit in {"万元", "万"}:
        return value * 10_000
    return value


def amount_label(prefix: str) -> str:
    prefix = re.sub(r"\s+", "", prefix)[-28:]
    if "合同金额" in prefix or "合同价款" in prefix or "签约金额" in prefix:
        return "合同金额"
    if re.search(r"成交价|成交金额|中标价|中标金额|投标报价", prefix):
        return "成交/中标报价"
    if "预算" in prefix:
        return "预算"
    if "限价" in prefix or "控制价" in prefix:
        return "限价/控制价"
    if "预估" in prefix and "金额" in prefix:
        return "预估金额"
    if "成本" in prefix and re.search(r"节约|节省|降低", prefix):
        return "效益/减排值（非报价）"
    if "成本" in prefix:
        return "成本"
    if "单价" in prefix:
        return "单价"
    if "报价" in prefix:
        return "报价"
    return "金额（标签待核）"


def amount_mentions(text: str) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    seen: set[tuple[str, float, str]] = set()
    for match in AMOUNT_RE.finditer(text):
        prefix = re.split(r"[。；|]", match.group("prefix"))[-1][-140:]
        if not (match.group("currency") or re.search(r"金额|价|预算|成本|费用", prefix)):
            continue
        if re.search(
            r"(?:金额|报价|预算|限价|成本|费用|价)\s*"
            r"(?:<=|>=|≤|≥|<|>|不超过|不低于|达到|每)\s*$",
            re.sub(r"\s+", "", prefix)[-24:],
        ) or re.search(r"(?:<=|>=|≤|≥|<|>)\s*$", prefix[-12:]):
            continue
        if re.search(
            r"采购文件.{0,10}(?:费用|售价)|手机CA|平台服务费|代理服务费|报名费",
            prefix[-32:],
        ):
            continue
        raw_number = match.group("num")
        unit = match.group("unit")
        value = normalize_amount(raw_number, unit)
        if value < 1 or value > 100_000_000_000:
            continue
        label = amount_label(prefix)
        if label == "金额（标签待核）" and not re.search(
            r"项目金额|总金额|估算金额|(?:含税|不含税)金额", prefix[-32:]
        ):
            continue
        key = (label, round(value, 2), unit)
        if key in seen:
            continue
        seen.add(key)
        compact = re.sub(r"\s+", "", match.group(0))
        result.append(
            {
                "label": label,
                "value": round(value, 2),
                "display": compact[-55:],
                "raw": f"{raw_number}{unit}",
            }
        )
        if len(result) >= 24:
            break
    return result


def detect_template_flags(text: str, entity_name: str) -> list[str]:
    flags: list[str] = []
    for sentence in sentence_candidates(text):
        if (
            re.search(r"建议.*模板|按照模板|此处填写|示例文本|待补充内容", sentence)
            and not re.search(
                r"导入模板|下载模板|批量导入|按照模板内容填写|按照模板要求填写",
                sentence,
            )
        ):
            flags.append(f"模板/示例措辞：{sentence[:140]}")
            if len(flags) >= 3:
                return flags
    normalized_entity = re.sub(r"\W+", "", entity_name)
    for name in PROJECT_NAME_RE.findall(text):
        normalized_name = re.sub(r"\W+", "", name)
        if len(normalized_name) < 4:
            continue
        if re.search(r"^(?:见|详见)|供应商须知|前附表|以采购文件为准", name):
            continue
        if normalized_name in normalized_entity or normalized_entity in normalized_name:
            continue
        overlap = len(set(normalized_name) & set(normalized_entity)) / max(
            1, len(set(normalized_name))
        )
        if overlap < 0.45:
            flags.append(f"正文出现其他项目名称，需核对是否模板残留：{name[:100]}")
        if len(flags) >= 3:
            break
    return flags


def evidence_for_entity(
    entity: dict[str, str], rows: list[dict[str, str]]
) -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    ranked = sorted(
        rows,
        key=lambda row: source_score(
            row, entity["entity_type"], entity["entity_name"]
        ),
        reverse=True,
    )
    evidence: list[dict[str, Any]] = []
    for row in ranked:
        if len(evidence) >= 10:
            break
        path = SOURCE_ROOT / row["relative_path"]
        if not path.exists():
            evidence.append({"source": row, "error": "源文件不存在", "facts": []})
            continue
        cache_path = CACHE_DIR / f"{row['sha256']}.json"
        if cache_path.exists():
            cached = json.loads(cache_path.read_text(encoding="utf-8"))
            text, error = cached["text"], cached["error"]
        else:
            text, error = extract_text(path)
            write_text(
                cache_path,
                json.dumps({"text": text, "error": error}, ensure_ascii=False),
            )
        if not text:
            continue
        evidence.append(
            {
                "source": row,
                "text": text,
                "error": error,
                "facts": extract_facts(text, entity["entity_name"]),
                "parties": party_mentions(text),
                "amounts": amount_mentions(text),
                "template_flags": detect_template_flags(text, entity["entity_name"]),
            }
        )
    return evidence, ranked


def infer_bid_status(
    entity: dict[str, str], rows: list[dict[str, str]]
) -> str:
    names = " ".join(row["filename"] for row in rows)
    if "弃标" in names:
        return "已弃标"
    if "废标公告" in names:
        return "已废标"
    if re.search(
        r"中标通知书|成交通知书|成交结果公告|成交公告|中选结果|中选通知书",
        names,
    ):
        return "已中标"
    if "13-中标" in entity["source_top_directories"]:
        return "已中标"
    if "04-投标中" in entity["source_top_directories"]:
        return "已投标"
    return "状态待核"


def infer_evidence_grade(
    entity: dict[str, str], rows: list[dict[str, str]], bid_status: str
) -> str:
    filenames = [row["filename"] for row in rows]
    paths = [row["relative_path"] for row in rows]
    if bid_status == "已中标":
        if any(
            re.search(
                r"中标通知书|成交通知书|成交结果公告|成交公告|中选结果|中选通知书",
                name,
            )
            for name in filenames
        ) or any("合同/收入/" in path for path in paths):
            return "strong"
        if "13-中标" in entity["source_top_directories"]:
            return "directory-only"
    if bid_status == "已废标" and any("废标公告" in name for name in filenames):
        return "strong"
    if bid_status == "已弃标" and any("弃标" in name for name in filenames):
        return "strong"
    if bid_status == "已投标" and any(
        re.search(r"投标|响应文件", name) for name in filenames
    ):
        return "strong"
    if bid_status == "已投标未中标" and any(
        re.search(r"未中标|落标|未成交", name) for name in filenames
    ):
        return "strong"
    if bid_status == "未投标" and any(
        re.search(r"未投标|不参与投标|放弃投标", name) for name in filenames
    ):
        return "strong"
    return "unknown"


def customer_from(parties: list[tuple[str, str]]) -> str:
    priorities = ["客户单位", "采购人", "采购单位", "建设单位", "使用单位", "甲方"]
    for role in priorities:
        values = [value for current_role, value in parties if current_role == role]
        if values:
            return Counter(values).most_common(1)[0][0]
    return ""


def period_from(evidence: list[dict[str, Any]]) -> str:
    patterns = (
        r"(?:服务期|服务期限|合同有效期|合作期限|工期|建设工期|交付期)"
        r"[^。；\n]{0,60}?\d+\s*(?:个?(?:日历)?日|个月|年)",
        r"合同(?:签订|生效)[^。；\n]{0,60}?\d+\s*(?:个?(?:日历)?日|个月|年)",
    )
    for item in evidence:
        text = item.get("text", "")
        for pattern in patterns:
            match = re.search(pattern, text)
            if match:
                value = re.sub(r"\s+", " ", match.group(0)).strip()
                if 6 <= len(value) <= 140:
                    return value
    return ""


def primary_amount(
    amount_evidence: list[tuple[str, dict[str, Any]]], entity_type: str
) -> tuple[float | None, str]:
    if entity_type in {"standard", "solution"}:
        return None, ""
    priorities = [
        "合同金额",
        "成交/中标报价",
        "预算",
        "限价/控制价",
        "预估金额",
        "报价",
        "成本",
    ]
    for label in priorities:
        values = [
            amount
            for _, amount in amount_evidence
            if amount["label"] == label
        ]
        if values:
            selected = Counter(round(item["value"], 2) for item in values).most_common(1)[0][0]
            return selected, label
    return None, ""


def money(value: float | None) -> str:
    if value is None:
        return "状态待核"
    return f"¥{value:,.2f}"


def source_rows_markdown(rows: list[dict[str, str]], primary_paths: set[str]) -> str:
    output = [
        "| 标记 | doc_id | 相对源根路径 | SHA-256 | 状态 |",
        "|---|---|---|---|---|",
    ]
    for row in rows:
        marker = "A" if row["source_pool"] == "A" else "C 保留"
        doc_id = f"SRC-{row['sha256'][:8]}"
        status = "主证据" if row["relative_path"] in primary_paths else "已聚合"
        path = row["relative_path"].replace("`", "'")
        output.append(
            f"| {marker} | {doc_id} | `{path}` | {row['sha256']} | {status} |"
        )
    return "\n".join(output)


def related_party_table(
    evidence: list[dict[str, Any]],
    rows: list[dict[str, str]],
    entity: dict[str, str],
) -> str:
    grouped: dict[tuple[str, str], set[str]] = defaultdict(set)
    quote_source_names = [
        row["filename"] for row in rows if "报价" in row["filename"]
    ]
    for item in evidence:
        source_name = item["source"]["filename"]
        for role, value in item.get("parties", []):
            grouped[(role, value)].add(source_name)
    if entity["entity_type"] == "solution":
        for row in rows:
            parts = row["relative_path"].split("/")
            if len(parts) >= 3 and parts[1] == "外部解决方案":
                vendor = parts[2].split("-", 1)[0].strip()
                grouped[("外部方案提供方（目录依据）", vendor)].add(row["filename"])
    quote_companies: dict[str, set[str]] = defaultdict(set)
    for name in quote_source_names:
        stem = Path(name).stem
        for pattern in (
            r"^([\u4e00-\u9fff]{2,12})(?:安全产品)?报价表",
            r"报价[-_ ]([\u4e00-\u9fff]{2,12})$",
        ):
            match = re.search(pattern, stem)
            if match:
                quote_companies[match.group(1)].add(name)
                break
    lines = [
        "| 相关方 | 卡片角色 | 依据 |",
        "|---|---|---|",
    ]
    for (role, value), names in sorted(grouped.items(), key=lambda item: (item[0][0], item[0][1])):
        if any("报价" in name for name in names) and role in {"供应商", "投标人", "响应人"}:
            card_role = "报价参考方（不据此认定为合作方）"
        else:
            card_role = role
        lines.append(f"| {value} | {card_role} | {'、'.join(sorted(names)[:3])} |")
    for value, names in sorted(quote_companies.items()):
        lines.append(
            f"| {value} | 报价方/报价参考方，不自动认定为合作方 | {'、'.join(sorted(names)[:3])} |"
        )
    if len(lines) == 2:
        lines.append("| 状态待核 | 现有主体文件的明确字段不足以确认 | — |")
    lines.extend(
        [
            "",
            "> [!warning] 报价方与合作方",
            "> 报价表、询价反馈或方案厂商中出现的主体只记为报价/能力参考方；"
            "只有生态合作记录、合同或明确的合作字段才能认定为合作方。",
        ]
    )
    return "\n".join(lines)


def evidence_boundary(entity: dict[str, str], parsed: int, total: int) -> str:
    return (
        f"- 本卡按 `{entity['entity_id']}` 聚合 {total} 个主体文件，"
        f"本次正文抽取覆盖 {parsed} 个可解析来源文件。"
        "\n- 客户、报价、成交、合同与服务期仅按明文或文件名中的明确状态记录；"
        "未能由原件字段确认的内容统一写“状态待核”。"
        "\n- 目录名只用于业务/归档状态线索，不等同于合同、验收或回款结论。"
        "\n- 扫描件、旧版 Office、压缩包、代码和媒体文件保留来源、SHA-256 与文件状态；"
        "未提取正文不等于原件无效。"
    )


def render_card(
    entity: dict[str, str],
    rows: list[dict[str, str]],
    evidence: list[dict[str, Any]],
    path: Path,
    preserve_sample: bool,
) -> dict[str, Any]:
    if preserve_sample:
        card = path.read_text(encoding="utf-8")
        source_count_match = re.search(r"^source_count:\s*(\d+)\s*$", card, re.MULTILINE)
        manifest = {
            "entity_id": entity["entity_id"],
            "entity_type": entity["entity_type"],
            "entity_name": entity["entity_name"],
            "card_path": str(path.relative_to(OUTPUT_DIR)),
            "source_count": len(rows),
            "parsed_count": "",
            "sample_preserved": "true",
            "amount_candidates": "",
            "amount_conflict": "",
            "template_flags": "",
            "extraction_error": "",
            "yaml_source_count": source_count_match.group(1) if source_count_match else "",
        }
        return manifest

    all_parties = [party for item in evidence for party in item.get("parties", [])]
    customer = customer_from(all_parties)
    bid_status = infer_bid_status(entity, rows) if entity["entity_type"] in {
        "project",
        "opportunity",
    } else "不适用"
    evidence_grade = infer_evidence_grade(entity, rows, bid_status)
    amount_evidence = [
        (item["source"]["filename"], amount)
        for item in evidence
        for amount in item.get("amounts", [])
    ]
    if bid_status in {"已弃标", "已废标"}:
        amount_evidence = [
            (filename, amount)
            for filename, amount in amount_evidence
            if amount["label"] != "成交/中标报价"
        ]
    chosen_amount, amount_label_text = primary_amount(amount_evidence, entity["entity_type"])
    amount_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for filename, amount in amount_evidence:
        amount_groups[amount["label"]].append({**amount, "source": filename})
    conflicts: list[str] = []
    for label, values in amount_groups.items():
        unique_values = sorted({round(value["value"], 2) for value in values})
        if len(unique_values) > 1 and label not in {"金额（标签待核）", "报价"}:
            conflicts.append(
                f"{label}存在不同数值："
                + "；".join(
                    f"{money(item['value'])}（{item['source']}，{item['raw']}）"
                    for item in values[:6]
                )
            )

    template_flags = [
        f"{item['source']['filename']}：{flag}"
        for item in evidence
        for flag in item.get("template_flags", [])
    ]
    parsed_count = len(evidence)
    period = period_from(evidence)
    folder, _, label = TYPE_INFO[entity["entity_type"]]
    business_status = entity["business_status"]
    if entity["entity_type"] == "project" and business_status == "已结束":
        business_text = "`已结束`（依据：来源位于 `07-已归档`；不等于验收/回款完成）"
    else:
        business_text = f"`{business_status}`"

    fact_lines: list[str] = []
    for item in evidence:
        for fact in item.get("facts", []):
            fact_lines.append(f"- **{item['source']['filename']}**：{fact}")
            if len(fact_lines) >= 8:
                break
        if len(fact_lines) >= 8:
            break
    if not fact_lines:
        fact_lines.append("- 当前候选主体文件未提取到可用于归卡的明确正文；保留文件名、路径与 SHA 供核验。")

    amount_lines = [
        "| 口径 | 数值 | 来源文件 | 原文片段 |",
        "|---|---:|---|---|",
    ]
    if amount_evidence:
        for filename, amount in amount_evidence[:24]:
            amount_lines.append(
                f"| {amount['label']} | {money(amount['value'])} | {filename.replace('|', '/')} | "
                f"`{amount['display'].replace('|', '/')}` |"
            )
    else:
        amount_lines.append("| 状态待核 | — | 未提取到带金额上下文的明文 | — |")

    risk_lines = [
        "- 来源状态只覆盖归档、投标、结果文件等现有证据；合同签署、验收、回款等后续状态未由本卡补齐。"
    ]
    if conflicts:
        risk_lines.append(
            "- **数字冲突/口径差异待核**：" + "；".join(conflicts[:4])
        )
    else:
        risk_lines.append(
            "- 未发现同字段下可自动判定的金额数值冲突；预算、限价、报价、成交价与合同额仍是不同口径，不能互相替代。"
        )
    if template_flags:
        risk_lines.append(
            "- **模板残留待核**：" + "；".join(template_flags[:5])
        )
    else:
        risk_lines.append("- 本次抽取未命中模板/示例措辞或其他项目名称规则；不代表原件绝无残留。")
    parsed_names = {item["source"]["relative_path"] for item in evidence}
    unparsed = [row for row in rows if row["relative_path"] not in parsed_names]
    if unparsed:
        extensions = Counter(row["extension"].casefold() for row in unparsed)
        summary = "、".join(f"{ext or '[无扩展名]'} {count}" for ext, count in extensions.most_common())
        risk_lines.append(
            f"- {len(unparsed)} 个来源未进入正文抽取，格式分布：{summary}；"
            "这些文件未据此做事实推断。"
        )

    related = related_party_table(evidence, rows, entity)
    sources = source_rows_markdown(rows, parsed_names)
    evidence_section = evidence_boundary(entity, parsed_count, len(rows))
    if entity["entity_type"] == "solution":
        ecosystem = (
            "\n## 生态相关方共享资料\n\n"
            "以下 3 个文件按所有者决定作为方案卡共享相关方资料，不另建实体，"
            "不增加 27 张方案卡的数量，也不在卡片中复制联系人手机号或邮箱。\n\n"
            "| doc_id | 相对源根路径 | 用途 |\n|---|---|---|\n"
            "| SRC-d65daf44 | `02-生态库/00-生态资源.xlsx` | 生态伙伴与合作记录 |\n"
            "| SRC-eb3365e5 | `02-生态库/02-11one平台联系人.xlsx` | 行业平台与专业公司协同 |\n"
            "| SRC-8967c7e2 | `02-生态库/安全厂家接口人.xlsx` | 安全产品询源与实施资源 |\n"
        )
    else:
        ecosystem = ""

    conclusion = (
        f"本卡聚合 {len(rows)} 个主体文件；"
        f"现有明文可支持 {len(evidence)} 个文件的部分事实核验，"
        f"业务状态记为 {business_status}，投标/成交状态记为 {bid_status}。"
        "客户、金额、期限或结果中未由原件明确支持的内容均保持待核。"
    )
    if entity["entity_type"] in {"solution", "standard"}:
        conclusion = (
            f"本卡聚合 {len(rows)} 个主体文件，用于沉淀{label}的类型、能力、流程或模板入口；"
            "现有原件未形成项目成交承诺，金额和适用项目边界以表内证据为准。"
        )

    amount_yaml = ""
    if chosen_amount is not None:
        amount_yaml = f"{chosen_amount:.2f}".rstrip("0").rstrip(".")
    card = f"""---
title: "{entity['entity_name'].replace('"', "'")}"
created: {RUN_DATE}
updated: {RUN_DATE}
type: {entity['entity_type']}
entity_id: "{entity['entity_id']}"
auto_generated: true
sample: {"true" if entity["entity_id"] in SAMPLE_CARDS else "false"}
status: complete
business_status: {business_status}
bid_status: {bid_status}
evidence_grade: {evidence_grade}
customer: "{customer.replace('"', "'")}"
amount: {amount_yaml}
period: "{period.replace('"', "'")}"
source_count: {len(rows)}
---

# {entity['entity_name']}

> [!info] 页面口径
> 页面类型：**{label}**。本卡按 `{entity['entity_id']}` 聚合 {len(rows)} 个主体文件，
> 不按目录或文件拆卡。来源目录只作状态线索；正文未支持的内容不猜、不补。

## 事实速览

| 字段 | 内容 |
|---|---|
| 客户/需求方 | {customer or "状态待核"} |
| 业务状态 | {business_text} |
| 投标/成交状态 | `{bid_status}` |
| 证据强度 | `{evidence_grade}` |
| 金额 | {money(chosen_amount)}{f"（{amount_label_text}）" if amount_label_text else ""} |
| 主体文件 | {len(rows)} 个（A 池 {sum(row['source_pool'] == 'A' for row in rows)}；C 保留 {sum(row['source_pool'] == 'C' for row in rows)}） |
| 来源目录 | {entity['source_top_directories'] or "状态待核"} |

## 一句话结论

{conclusion}

## 关键事实

{chr(10).join(fact_lines)}

## 金额来源与数字口径

{chr(10).join(amount_lines)}

> 金额字段优先取合同额，其次取成交/中标报价，再次保留预算、限价与普通报价；
> 不同口径不互相覆盖，也不能用页面中的汇总值替代原件。

## 相关方

{related}
{ecosystem}
## 风险与待办

{chr(10).join(risk_lines)}

## 证据边界

{evidence_section}

## 来源

{sources}

## 相关

- [[../index|总索引]]
"""
    write_text(path, card)
    return {
        "entity_id": entity["entity_id"],
        "entity_type": entity["entity_type"],
        "entity_name": entity["entity_name"],
        "card_path": str(path.relative_to(OUTPUT_DIR)),
        "source_count": len(rows),
        "parsed_count": parsed_count,
        "sample_preserved": "false",
        "amount_candidates": len(amount_evidence),
        "amount_conflict": len(conflicts),
        "template_flags": len(template_flags),
        "extraction_error": "",
        "yaml_source_count": len(rows),
    }


def batch_qa_markdown(
    batch_id: str,
    title: str,
    types: tuple[str, ...],
    entities: list[dict[str, str]],
    manifest: list[dict[str, Any]],
) -> str:
    counts = Counter("total" for _ in entities)
    type_counts = Counter(entity["entity_type"] for entity in entities)
    source_total = sum(int(row["source_count"]) for row in manifest)
    parsed_total = sum(
        int(row["parsed_count"]) for row in manifest if str(row["parsed_count"]).isdigit()
    )
    conflicts = sum(
        int(row["amount_conflict"]) for row in manifest if str(row["amount_conflict"]).isdigit()
    )
    template_flags = sum(
        int(row["template_flags"]) for row in manifest if str(row["template_flags"]).isdigit()
    )
    preserved = sum(row["sample_preserved"] == "true" for row in manifest)
    source_mismatch = [
        row
        for row in manifest
        if row["sample_preserved"] != "true"
        and str(row["source_count"]) != str(row["yaml_source_count"])
    ]
    type_lines = "\n".join(
        f"| {TYPE_INFO[type_name][0]} | {type_counts[type_name]} |"
        for type_name in types
    )
    missing = [
        entity["entity_id"]
        for entity in entities
        if not find_card(entity).exists()
    ]
    return f"""# 批次 {batch_id} 核验：{title}

- 卡数：**{counts['total']}**
- 实体来源文件：**{source_total:,}**
- 正文可解析来源：**{parsed_total:,}**
- 保留的已批准样本：**{preserved}**
- 金额口径冲突：**{conflicts}**
- 模板/其他项目名提示：**{template_flags}**
- 缺失卡片：**{len(missing)}**
- YAML `source_count` 不一致：**{len(source_mismatch)}**

| 类型 | 卡数 |
|---|---:|
{type_lines}

## 验收口径

- 每张卡对应一个 `entity_id`，同一实体全部主体文件聚合到一张卡。
- 来源表逐行保留 A/C 状态、`doc_id`、相对路径和 SHA-256。
- 客户、金额、报价方、合作方均按明确字段抽取；报价文件中的主体不被自动认定为合作方。
- 数字冲突、模板/其他项目名提示和未解析格式均写入对应卡片风险段。
- 本批 `manifest.csv` 是可机器复核的逐卡台账。
"""


def run_batch(batch_id: str) -> None:
    types, title = BATCHES[batch_id]
    entities_all = read_csv(INVENTORY / "实体清单.csv")
    source_rows = read_csv(INVENTORY / "主体文件台账.csv")
    by_entity: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in source_rows:
        by_entity[row["entity_id"]].append(row)
    entities = [row for row in entities_all if row["entity_type"] in types]

    def process_entity(
        indexed_entity: tuple[int, dict[str, str]],
    ) -> tuple[dict[str, Any], str]:
        index, entity = indexed_entity
        rows = by_entity[entity["entity_id"]]
        evidence, _ = evidence_for_entity(entity, rows)
        path = find_card(entity)
        preserve = (
            entity["entity_id"] in SAMPLE_CARDS
            and SAMPLE_CARDS[entity["entity_id"]].exists()
        )
        card_manifest = render_card(entity, rows, evidence, path, preserve)
        message = (
            f"[{batch_id}] {index:02d}/{len(entities):02d} "
            f"{entity['entity_id']} {entity['entity_name'][:42]} "
            f"source={len(rows)} parsed={card_manifest['parsed_count'] or '-'}"
        )
        return card_manifest, message

    with ThreadPoolExecutor(max_workers=WORKERS) as executor:
        results = list(executor.map(process_entity, enumerate(entities, 1)))
    manifest = [item[0] for item in results]
    for _, message in results:
        print(message)

    prefix = f"batch{batch_id}"
    write_csv(
        BATCH_DIR / f"{prefix}-manifest.csv",
        manifest,
        [
            "entity_id",
            "entity_type",
            "entity_name",
            "card_path",
            "source_count",
            "parsed_count",
            "sample_preserved",
            "amount_candidates",
            "amount_conflict",
            "template_flags",
            "extraction_error",
            "yaml_source_count",
        ],
    )
    write_text(
        BATCH_DIR / f"{prefix}-qa.md",
        batch_qa_markdown(batch_id, title, types, entities, manifest),
    )
    write_text(
        BATCH_DIR / f"{prefix}-done.json",
        json.dumps(
            {
                "batch": int(batch_id),
                "title": title,
                "completed_at": datetime.now().astimezone().isoformat(),
                "entities": len(entities),
                "sources": sum(int(row["source_count"]) for row in manifest),
                "cards": len(manifest),
            },
            ensure_ascii=False,
            indent=2,
        ),
    )


def build_index() -> None:
    entities = read_csv(INVENTORY / "实体清单.csv")
    by_type: dict[str, list[dict[str, str]]] = defaultdict(list)
    for entity in entities:
        by_type[entity["entity_type"]].append(entity)
    manifest_rows: list[dict[str, str]] = []
    for path in sorted(BATCH_DIR.glob("batch*-manifest.csv")):
        manifest_rows.extend(read_csv(path))
    manifest_by_id = {row["entity_id"]: row for row in manifest_rows}
    sections: list[str] = []
    for type_name in ["project", "opportunity", "solution", "standard", "procurement"]:
        folder, _, _ = TYPE_INFO[type_name]
        lines = [
            f"## {folder}（{len(by_type[type_name])}）",
            "",
            "| entity_id | 实体 | 来源数 |",
            "|---|---|---:|",
        ]
        for entity in by_type[type_name]:
            target = find_card(entity)
            relative = target.relative_to(VAULT).with_suffix("").as_posix()
            lines.append(
                f"| `{entity['entity_id']}` | [[{relative}\\|{entity['entity_name']}]] | "
                f"{manifest_by_id.get(entity['entity_id'], {}).get('source_count', '')} |"
            )
        sections.append("\n".join(lines))
    type_counts = Counter(entity["entity_type"] for entity in entities)
    summary_rows = "\n".join(
        f"| {TYPE_INFO[type_name][0]} | {type_counts[type_name]} |"
        for type_name in ["project", "opportunity", "solution", "standard", "procurement"]
    )
    index = f"""# {OUTPUT_DIR.name}实体卡

 > [!success] 全量建卡完成
> 已生成 **{len(entities)} 张**实体卡。每批均有独立 QA 与 manifest。

| 类型 | 卡数 |
|---|---:|
{summary_rows}
| **合计** | **{len(entities)}** |

{chr(10).join(sections)}

## 批次验收

- [[../build/batches/batch1-qa|批次1：项目]]
- [[../build/batches/batch2-qa|批次2：机会]]
- [[../build/batches/batch3-qa|批次3：方案]]
- [[../build/batches/batch4-qa|批次4：规范与采购]]
- [[../build/全量建卡QA|全量 QA]]
"""
    write_text(VAULT / "index.md", index)


def build_final_qa() -> None:
    entities = read_csv(INVENTORY / "实体清单.csv")
    inventory_sources = read_csv(INVENTORY / "主体文件台账.csv")
    expected = Counter(entity["entity_type"] for entity in entities)
    cards = [find_card(entity) for entity in entities]
    missing = [str(path) for path in cards if not path.exists()]
    unique_paths = {path.resolve() for path in cards if path.exists()}
    duplicate_paths = len(cards) - len(unique_paths) - len(missing)
    manifests: list[dict[str, str]] = []
    for path in sorted(BATCH_DIR.glob("batch*-manifest.csv")):
        manifests.extend(read_csv(path))
    manifest_ids = {row["entity_id"] for row in manifests}
    expected_ids = {entity["entity_id"] for entity in entities}
    source_total = sum(int(row["source_count"]) for row in manifests)
    parsed_total = sum(
        int(row["parsed_count"]) for row in manifests if row["parsed_count"].isdigit()
    )
    conflicts = sum(
        int(row["amount_conflict"]) for row in manifests if row["amount_conflict"].isdigit()
    )
    template_flags = sum(
        int(row["template_flags"]) for row in manifests if row["template_flags"].isdigit()
    )
    type_lines = "\n".join(
        f"| {TYPE_INFO[type_name][0]} | {expected[type_name]} |"
        for type_name in ["project", "opportunity", "solution", "standard", "procurement"]
    )
    report = f"""# {OUTPUT_DIR.name}：{len(entities)} 张实体卡全量 QA

生成时间：{datetime.now().astimezone().isoformat()}

## 数量闭环

| 类型 | 目标卡数 |
|---|---:|
{type_lines}
| **合计** | **{len(entities)}** |

- 已找到卡片：**{len(cards) - len(missing)} / {len(entities)}**
- 缺失：**{len(missing)}**
- 重复卡片路径：**{duplicate_paths}**
- manifest 实体：**{len(manifest_ids)} / {len(expected_ids)}**
- 主体来源覆盖：**{source_total:,}**，与来源台账 {len(inventory_sources):,} 个的差额：**{source_total - len(inventory_sources):+,}**
- 正文可解析来源：**{parsed_total:,}**
- 金额同字段冲突：**{conflicts}**
- 模板/其他项目名提示：**{template_flags}**

## 五项质量标准落点

| 标准 | 卡内位置 | 校验 |
|---|---|---|
| 数字冲突要标 | `金额来源与数字口径`、`风险与待办` | 同口径不同数值单列；不同金额口径不混用 |
| 金额来源要说清 | 金额表逐行保留数值、文件名与原文片段 | 无金额时写状态待核 |
| 模板残留要标 | `风险与待办` | 命中模板措辞或其他项目名时提示待核 |
| 报价方≠合作方 | `相关方`、固定 warning | 报价文件主体不自动认定合作方 |
| 不猜不补不圆场 | `事实速览`、`证据边界` | 无明文客户/金额/结果均待核 |

## 完整性说明

- {len(entities)} 张卡按实体聚合，不按来源文件拆卡。
- 所有可追溯来源均保留在 manifest 与来源台账中。
- 未删除、移动或改写任何业务原件。
"""
    if missing:
        report += "\n## 缺失卡片\n\n" + "\n".join(f"- `{path}`" for path in missing) + "\n"
    write_text(BUILD / "全量建卡QA.md", report)
    write_text(
        BUILD / "全量建卡总结.json",
        json.dumps(
            {
                "generated_at": datetime.now().astimezone().isoformat(),
                "targets": dict(expected),
                "target_total": len(entities),
                "cards_found": len(cards) - len(missing),
                "missing": len(missing),
                "source_total": source_total,
                "parsed_total": parsed_total,
                "amount_conflicts": conflicts,
                "template_flags": template_flags,
            },
            ensure_ascii=False,
            indent=2,
        ),
    )


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Build entity cards from an inventory and source evidence."
    )
    parser.add_argument(
        "--source-dir", type=Path, required=True, help="原始语料根目录"
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        required=True,
        help="输出目录（需包含 entity-inventory 预清单）",
    )
    parser.add_argument(
        "--batch",
        choices=[*sorted(BATCHES), "all"],
        default="all",
        help="只运行指定批次，默认 all",
    )
    parser.add_argument(
        "--workers",
        type=int,
        default=None,
        help="并发生成实体卡的线程数，默认 CPU 核数",
    )
    parser.add_argument(
        "--run-date",
        default=None,
        help="写入卡片的日期（YYYY-MM-DD），默认今天",
    )
    args = parser.parse_args()

    global SOURCE_ROOT, OUTPUT_DIR, VAULT, INVENTORY, BUILD, BATCH_DIR
    global CACHE_DIR, RUN_DATE, WORKERS
    SOURCE_ROOT = args.source_dir.resolve()
    OUTPUT_DIR = args.output_dir.resolve()
    VAULT = OUTPUT_DIR / "vault"
    INVENTORY = OUTPUT_DIR / "entity-inventory"
    BUILD = OUTPUT_DIR / "build"
    BATCH_DIR = BUILD / "batches"
    CACHE_DIR = BUILD / ".text-cache"
    RUN_DATE = args.run_date or datetime.now().astimezone().date().isoformat()
    WORKERS = args.workers or max(1, os.cpu_count() or 1)

    if not SOURCE_ROOT.is_dir():
        parser.error(f"source directory not found: {SOURCE_ROOT}")
    required_inventory = (INVENTORY / "实体清单.csv", INVENTORY / "主体文件台账.csv")
    missing_inventory = [str(path) for path in required_inventory if not path.is_file()]
    if missing_inventory:
        parser.error("missing inventory files: " + ", ".join(missing_inventory))
    if WORKERS < 1:
        parser.error("--workers must be at least 1")

    VAULT.mkdir(parents=True, exist_ok=True)
    BATCH_DIR.mkdir(parents=True, exist_ok=True)
    batch_ids = sorted(BATCHES) if args.batch == "all" else [args.batch]
    for batch_id in batch_ids:
        run_batch(batch_id)
    if "4" in batch_ids:
        build_index()
        build_final_qa()


if __name__ == "__main__":
    main()
